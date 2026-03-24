# /// script
# dependencies = [
#   "python-pptx",
# ]
# ///

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
import os
from pathlib import Path

prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)

# Get the directory containing the visualizations (relative to script location)
script_dir = Path(__file__).parent
# Go up to project root (from scripts/analysis/ to project root)
project_root = script_dir.parent.parent
viz_dir = project_root / "visualization" / "viz"

# Verify the directory exists
if not viz_dir.exists():
    print(f"Error: Visualization directory not found at {viz_dir}")
    exit(1)

# Get all unique level names (smXXX_true format)
level_files = sorted(viz_dir.glob("stimuli_sm*_true_agent2_experienced1_and_agent3_experienced1.png"))
levels = [f.stem.replace("stimuli_", "").replace("_agent2_experienced1_and_agent3_experienced1", "") for f in level_files]

print(f"Found {len(levels)} levels to process")

for level in levels:
    # Add a blank slide
    blank_slide_layout = prs.slide_layouts[6]  # Blank layout
    slide = prs.slides.add_slide(blank_slide_layout)
    
    # Add title at the top
    title_box = slide.shapes.add_textbox(
        Inches(0.5), 
        Inches(0.3), 
        Inches(9), 
        Inches(0.6)
    )
    title_frame = title_box.text_frame
    title_frame.text = level.upper()
    
    # Format title
    title_paragraph = title_frame.paragraphs[0]
    title_paragraph.alignment = PP_ALIGN.CENTER
    title_paragraph.font.size = Pt(32)
    title_paragraph.font.bold = True
    
    # Add first image (experienced1) - left side
    img1_path = viz_dir / f"stimuli_{level}_agent2_experienced1_and_agent3_experienced1.png"
    if img1_path.exists():
        left = Inches(0.5)
        top = Inches(1.2)
        width = Inches(4.25)
        slide.shapes.add_picture(str(img1_path), left, top, width=width)
    else:
        print(f"  Warning: Image not found: {img1_path}")
    
    # Add second image (experienced2) - right side with spacing
    img2_path = viz_dir / f"stimuli_{level}_agent2_experienced2_and_agent3_experienced2.png"
    if img2_path.exists():
        left = Inches(5.25)
        top = Inches(1.2)
        width = Inches(4.25)
        slide.shapes.add_picture(str(img2_path), left, top, width=width)
    else:
        print(f"  Warning: Image not found: {img2_path}")
    
    print(f"Added slide for {level}")

# Save the presentation (relative to script location)
output_path = script_dir / "path_comparison_presentation.pptx"
prs.save(str(output_path))
print(f"\nPresentation saved to: {output_path}")
print(f"Total slides: {len(levels)}")

