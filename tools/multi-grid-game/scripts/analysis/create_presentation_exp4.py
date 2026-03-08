# /// script
# dependencies = [
#   "python-pptx",
# ]
# ///

import os
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)

# Get the directory containing the visualizations (relative to script location)
script_dir = Path(__file__).parent
# Go up to project root (from scripts/analysis/ to project root)
project_root = script_dir.parent.parent
viz_dir = project_root / "visualization" / "viz_exp4"

# Verify the directory exists
if not viz_dir.exists():
    print(f"Error: Visualization directory not found at {viz_dir}")
    exit(1)

# Get all unique level names (smXXX_exp4 format)
level_files = sorted(
    viz_dir.glob("stimuli_sm*_exp4_agent1_experienced1_and_agent2_experienced1.png")
)
levels = [
    f.stem.replace("stimuli_", "").replace(
        "_agent1_experienced1_and_agent2_experienced1", ""
    )
    for f in level_files
]

print(f"Found {len(levels)} levels to process")

for level in levels:
    # Add a blank slide
    blank_slide_layout = prs.slide_layouts[6]  # Blank layout
    slide = prs.slides.add_slide(blank_slide_layout)

    # Add title at the top
    title_box = slide.shapes.add_textbox(
        Inches(0.5), Inches(0.3), Inches(9), Inches(0.6)
    )
    title_frame = title_box.text_frame
    title_frame.text = level.upper()

    # Format title
    title_paragraph = title_frame.paragraphs[0]
    title_paragraph.alignment = PP_ALIGN.CENTER
    title_paragraph.font.size = Pt(32)
    title_paragraph.font.bold = True

    # Calculate dimensions for 3 images in a row
    image_width = Inches(3)
    spacing = Inches(0.25)
    total_width = (image_width * 3) + (spacing * 2)
    start_left = (Inches(10) - total_width) / 2
    top = Inches(1.5)

    # Add first image (experienced1) - left
    img1_path = (
        viz_dir / f"stimuli_{level}_agent1_experienced1_and_agent2_experienced1.png"
    )
    if img1_path.exists():
        left = start_left
        slide.shapes.add_picture(str(img1_path), left, top, width=image_width)
    else:
        print(f"  Warning: Image not found: {img1_path}")

    # Add second image (experienced2) - middle
    img2_path = (
        viz_dir / f"stimuli_{level}_agent1_experienced2_and_agent2_experienced2.png"
    )
    if img2_path.exists():
        left = start_left + image_width + spacing
        slide.shapes.add_picture(str(img2_path), left, top, width=image_width)
    else:
        print(f"  Warning: Image not found: {img2_path}")

    # Add third image (experienced3) - right
    img3_path = (
        viz_dir / f"stimuli_{level}_agent1_experienced3_and_agent2_experienced3.png"
    )
    if img3_path.exists():
        left = start_left + (image_width + spacing) * 2
        slide.shapes.add_picture(str(img3_path), left, top, width=image_width)
    else:
        print(f"  Warning: Image not found: {img3_path}")

    print(f"Added slide for {level}")

# Save the presentation (relative to script location)
output_path = script_dir / "path_comparison_presentation_exp4.pptx"
prs.save(str(output_path))
print(f"\nPresentation saved to: {output_path}")
print(f"Total slides: {len(levels)}")
